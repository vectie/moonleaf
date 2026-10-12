#!/usr/bin/env python3
"""Exercise the built Node adapter; keep every stdout/stderr and all input bytes."""
from pathlib import Path
import hashlib
import importlib.metadata
import json
import subprocess
import sys
import zipfile
import xml.etree.ElementTree as ET

source = Path(__file__).resolve().parents[1]
build, evidence = map(Path, sys.argv[1:])
evidence.mkdir(parents=True, exist_ok=False)
cli = build / 'js/debug/build/cmd/xlsx-compare/xlsx-compare.js'
fixture = source / 'kit/workbook-compare'
original, copy, styled = [fixture/name for name in ['original.xlsx','client-copy.xlsx','style-only.xlsx']]
sha = lambda path: hashlib.sha256(path.read_bytes()).hexdigest()
provenance = json.loads((fixture/'PROVENANCE.json').read_text())
for name, digest in provenance['sha256'].items():
    assert sha(fixture/name) == digest
calls = []

def call(name, args, success=True, fmt='json'):
    argv = ['node', str(cli)] + [str(a) for a in args]
    if fmt is not None:
        argv.append(fmt)
    p = subprocess.run(argv, text=True, capture_output=True)
    (evidence/(name+'.stdout')).write_text(p.stdout)
    (evidence/(name+'.stderr')).write_text(p.stderr)
    calls.append({'name':name,'argv':argv,'exit_code':p.returncode,
                  'stdout_sha256':hashlib.sha256(p.stdout.encode()).hexdigest(),
                  'stderr_sha256':hashlib.sha256(p.stderr.encode()).hexdigest()})
    (evidence/'calls.json').write_text(json.dumps(calls, indent=2)+'\n')
    assert (p.returncode == 0) == success, (name, p.stderr)
    if not success:
        assert p.stdout == '' and p.stderr
        return p.stderr
    if fmt in ('json', None):
        return json.loads(p.stdout)
    return p.stdout

report = call('comparison', [original,copy])
assert report['schema'] == 'moonleaf.workbook-comparison.v1'
assert report['original_sha256'] == sha(original) and report['copy_sha256'] == sha(copy)
assert not report['identical_bytes'] and report['supported_content_changed']
comparison = report['comparison']
assert comparison['changed_cells'] == 6 and not comparison['date_system_changed']
assert comparison['before_sheets'] == ['Budget & 雪','Hidden notes','Old appendix']
assert comparison['after_sheets'] == ['Hidden notes','Budget & 雪','New appendix']
sheets = {s['name']:s for s in comparison['sheets']}
assert list(sheets) == ['Budget & 雪','Hidden notes','Old appendix','New appendix']
budget = sheets['Budget & 雪']
assert budget['changes'] == ['position','merges','cells']
assert budget['before_merges'] == ['A5:B5'] and budget['after_merges'] == ['A5:C5']
cells = {c['address']:c for c in budget['cells']}
assert list(cells) == ['B2','D2','E2','F2','G2','H2']
assert cells['B2']['before']['value'] == 'Client Alpha'
assert cells['B2']['after']['value'] == 'Client 雪😀 & <new>\nsecond line'
assert cells['D2']['changes'] == ['formula']
assert cells['D2']['before']['formula'] == 'C2*2' and cells['D2']['after']['formula'] == 'C2*3'
assert cells['D2']['before']['value'] == cells['D2']['after']['value'] == '25'
assert cells['E2']['changes'] == ['cached_value']
assert cells['E2']['before']['formula'] == cells['E2']['after']['formula'] == 'SUM(C2:C3)'
assert (cells['E2']['before']['value'],cells['E2']['after']['value']) == ('20','21')
assert cells['F2']['after'] is None and cells['F2']['changes'] == ['removed']
assert cells['H2']['before'] is None and cells['H2']['changes'] == ['added']
assert cells['H2']['after']['formula'] is None
assert cells['H2']['after']['value'] == '=literal text, not a formula'
assert cells['G2']['after']['value_type'] == 'boolean' and cells['G2']['after']['value'] == 'true'
assert sheets['Hidden notes']['changes'] == ['position','visibility']
assert sheets['Hidden notes']['before_state'] == 'hidden' and sheets['Hidden notes']['after_state'] == 'visible'
assert sheets['Old appendix']['after_position'] is None and sheets['Old appendix']['cells'] == []
assert sheets['New appendix']['before_position'] is None and sheets['New appendix']['cells'] == []
assert len(comparison['warnings']) == 4

text = call('readable', [original,copy], fmt='text')
assert 'D2 [formula]' in text and 'E2 [cached_value]' in text
assert 'stored number cache="25"' in text and 'cache freshness is unknown' in text
assert sha(original) in text and sha(copy) in text

same = call('same-file', [original,original], fmt=None)
assert same['identical_bytes'] and not same['supported_content_changed']
assert same['comparison']['changed_cells'] == 0 and same['comparison']['sheets'] == []
style = call('style-only', [original,styled])
assert not style['identical_bytes'] and not style['supported_content_changed']
assert style['comparison']['sheets'] == []
reverse = call('reverse', [copy,original])
reverse_cells = {c['address']:c for s in reverse['comparison']['sheets'] if s['name']=='Budget & 雪' for c in s['cells']}
for address, item in cells.items():
    assert reverse_cells[address]['before'] == item['after']
    assert reverse_cells[address]['after'] == item['before']

ns = {'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
def modified(name, part, transform):
    target = evidence/(name+'.xlsx')
    with zipfile.ZipFile(original) as src, zipfile.ZipFile(target,'w',zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            content = src.read(info.filename)
            if info.filename == part:
                content = transform(content)
            dst.writestr(info, content)
    return target

def set_epoch(data):
    root = ET.fromstring(data)
    pr = root.find('s:workbookPr',ns)
    assert pr is not None
    pr.set('date1904','1')
    return ET.tostring(root,encoding='utf-8')
epoch = modified('epoch-only','xl/workbook.xml',set_epoch)
epoch_report = call('epoch', [original,epoch])
assert epoch_report['supported_content_changed']
assert epoch_report['comparison']['date_system_changed'] and epoch_report['comparison']['changed_cells'] == 0
assert epoch_report['comparison']['sheets'] == []

def duplicate(data):
    root = ET.fromstring(data)
    row = root.find('s:sheetData/s:row',ns)
    row.append(ET.fromstring(ET.tostring(row.find('s:c',ns))))
    return ET.tostring(root,encoding='utf-8')
dup = modified('duplicate-cell','xl/worksheets/sheet1.xml',duplicate)
assert 'duplicate comparison cell' in call('duplicate-refused', [original,dup], success=False)
call('bad-mode',[original,copy],success=False,fmt='csv')
call('bad-args',[],success=False,fmt=None)
call('missing-file',[original,evidence/'absent.xlsx'],success=False)
call('directory-input',[original,evidence],success=False)
invalid = evidence/'not-office.xlsx'
invalid.write_text('Not an XLSX file')
call('invalid-archive',[original,invalid],success=False)
oversize = evidence/'over-limit.xlsx'
with oversize.open('wb') as f:
    f.truncate(8388609)
call('oversize-file',[original,oversize],success=False)

# A separate producer/reader verifies the retained fixture semantics and caches.
import openpyxl
assert importlib.metadata.version('openpyxl') == '3.1.5'
original_formulas = openpyxl.load_workbook(original,data_only=False)
copy_formulas = openpyxl.load_workbook(copy,data_only=False)
original_cache = openpyxl.load_workbook(original,data_only=True)
copy_cache = openpyxl.load_workbook(copy,data_only=True)
a,b = original_formulas['Budget & 雪'],copy_formulas['Budget & 雪']
assert a['D2'].value == '=C2*2' and b['D2'].value == '=C2*3'
assert original_cache['Budget & 雪']['D2'].value == copy_cache['Budget & 雪']['D2'].value == 25
assert original_cache['Budget & 雪']['E2'].value == 20 and copy_cache['Budget & 雪']['E2'].value == 21
assert b['H2'].data_type == 's' and b['H2'].value == '=literal text, not a formula'
assert b['A2'].value == '001'
assert original_formulas['Hidden notes'].sheet_state == 'hidden'
assert copy_formulas['Hidden notes'].sheet_state == 'visible'
for book in [original_formulas,copy_formulas,original_cache,copy_cache]:
    book.close()
for name,digest in provenance['sha256'].items():
    assert sha(fixture/name) == digest
result = {'status':'pass','cli_cases':len(calls),'changed_cells':comparison['changed_cells'],
          'independent_reader':'openpyxl 3.1.5','producer':'XlsxWriter 3.2.9',
          'retained_inputs_unchanged':True,'fixture_sha256':provenance['sha256'],
          'formula_and_cache_changes_independent':True,'style_only_not_content_equal_to_bytes':True,
          'scope':'Read-only supported-content comparison, no calculation or visual fidelity qualification.'}
(evidence/'RESULT.json').write_text(json.dumps(result,indent=2)+'\n')
print(json.dumps(result))
