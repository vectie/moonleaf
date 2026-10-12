"""Actual CLI acceptance; requires Node, Python 3 and retained openpyxl 3.1.5.
Usage: python3 tests/xlsx_named_cli.py ABSOLUTE_BUILD FRESH_EVIDENCE_DIRECTORY
"""
from pathlib import Path
import csv
import hashlib
import io
import json
import os
import subprocess
import sys
import zipfile
import xml.etree.ElementTree as ET
import openpyxl

source = Path(__file__).resolve().parents[1]
build = Path(sys.argv[1]).resolve()
out = Path(sys.argv[2]).resolve()
out.mkdir(parents=True, exist_ok=False)
cli = build / 'js/debug/build/cmd/xlsx-named-table/xlsx-named-table.js'
fixture = source / 'kit/xlsx-named-table/orders.xlsx'
query = source / 'kit/xlsx-named-table/query.json'
digest = lambda data: hashlib.sha256(data).hexdigest()
original = fixture.read_bytes()
calls = []


def call(*args, good=True):
    proc = subprocess.run(['node', str(cli), *map(str, args)], capture_output=True)
    number = len(calls) + 1
    (out / f'{number:02d}.stdout').write_bytes(proc.stdout)
    (out / f'{number:02d}.stderr').write_bytes(proc.stderr)
    calls.append({'arguments': list(map(str, args)), 'exit_code': proc.returncode, 'expected_success': good})
    (out / 'calls.json').write_text(json.dumps(calls, indent=2) + '\n')
    assert (proc.returncode == 0) == good, (args, proc.stdout, proc.stderr)
    if good:
        return json.loads(proc.stdout)
    assert proc.stdout == b'', (args, proc.stdout)


def export(name, fmt='json', recipe=query, workbook=fixture, table='Orders'):
    dest = out / name
    report = call('export', workbook, table, recipe, fmt, dest)
    assert report['source_sha256'] == digest(Path(workbook).read_bytes())
    assert report['output_sha256'] == digest(dest.read_bytes())
    assert report['readback_equal'] is True
    assert os.stat(dest).st_mode & 0o777 == 0o600
    return report, dest


ns = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}

def mutate(name, part, transform):
    with zipfile.ZipFile(io.BytesIO(original)) as z:
        entries = {key: z.read(key) for key in z.namelist()}
    entries[part] = transform(entries[part])
    path = out / name
    with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as z:
        for key, value in entries.items():
            z.writestr(key, value)
    return path


def refusal(name, workbook=fixture, table='Orders', recipe=query):
    dest = out / f'{name}.json'
    call('export', workbook, table, recipe, 'json', dest, good=False)
    assert not dest.exists()


catalog = call('inspect', fixture)
assert catalog['source_sha256'] == digest(original)
assert [t['name'] for t in catalog['tables']] == ['Orders', 'Regions', 'Feed']
assert [t['sheet_name'] for t in catalog['tables']] == ['Sales Q4', 'Ops', 'Hidden Feed']
assert catalog['tables'][0]['import_reference'] == 'B3:F6'
assert catalog['tables'][0]['data_rows'] == 3
assert catalog['tables'][2]['header_rows'] == 0
book = openpyxl.load_workbook(fixture, data_only=False)
for table in catalog['tables']:
    actual = book[table['sheet_name']].tables[table['name']]
    assert actual.ref == table['reference']
    assert [c.name for c in actual.tableColumns] == table['columns']
assert book['Sales Q4'].row_dimensions[6].hidden
assert book['Hidden Feed'].sheet_state == 'hidden'
assert book['Sales Q4']['D7'].data_type == 'f'

expected = [{'id': '003', 'amount': 20, 'note': 'priority\nship'},
            {'id': '001', 'amount': 12.5, 'note': '=SUM(1,2)'}]
report, data = export('orders.json')
assert json.loads(data.read_text()) == expected
assert report['worksheet_rows'] == [6, 4]
assert (report['input_rows'], report['matched_rows'], report['output_rows']) == (3, 2, 2)
assert report['query_sha256'] == digest(query.read_bytes())
report, data = export('orders.csv', 'csv')
assert list(csv.reader(io.StringIO(data.read_text()))) == [['id', 'amount', 'note'], ['003', '20', 'priority\nship'], ['001', '12.5', '=SUM(1,2)']]
report, data = export('orders.xlsx', 'xlsx')
exported = openpyxl.load_workbook(data, data_only=False)
rows = list(exported.active.values)
assert rows == [('id', 'amount', 'note'), ('003', 20, 'priority\nship'), ('001', 12.5, '=SUM(1,2)')]
assert exported.active['C3'].data_type == 's'
assert exported.active['B2'].data_type == 'n'
assert exported.active['A2'].data_type == 's'
with zipfile.ZipFile(data) as z:
    assert z.testzip() is None
    root = ET.fromstring(z.read('xl/worksheets/sheet1.xml'))
    assert not root.findall('.//s:f', ns)
report, data = export('all.json', recipe='-')
all_rows = json.loads(data.read_text())
assert [r['id'] for r in all_rows] == ['001', '002', '003']
assert [r['active'] for r in all_rows] == [True, False, True]
assert report['worksheet_rows'] == [4, 5, 6]
assert report['query_sha256'] is None
report, data = export('feed.json', recipe='-', table='Feed')
assert json.loads(data.read_text()) == [{'id': '101', 'amount': 3}, {'id': '102', 'amount': 4}]
assert report['worksheet_rows'] == [2, 3]

refusal('missing-name', table='orders')
invalid_query = out / 'invalid-query.json'
invalid_query.write_text('{}')
refusal('bad-query', recipe=invalid_query)
invalid_utf8 = out / 'invalid-utf8.json'
invalid_utf8.write_bytes(b'\xff')
refusal('bad-utf8', recipe=invalid_utf8)

part = 'xl/tables/table1.xml'
variants = {
    'header-mismatch': lambda b: b.replace(b'name="id"', b'name="other"'),
    'calculated-column': lambda b: b.replace(b'<tableColumn id="2" name="customer"/>', b'<tableColumn id="2" name="customer"><calculatedColumnFormula>1+1</calculatedColumnFormula></tableColumn>'),
    'bad-range': lambda b: b.replace(b'B3:F7', b'B3:Z7'),
    'query-table': lambda b: b.replace(b'<table ', b'<table tableType="queryTable" '),
    'duplicate-name': lambda b: b.replace(b'displayName="Orders"', b'displayName="Regions"'),
    'bad-totals': lambda b: b.replace(b'totalsRowCount="1"', b'totalsRowCount="2"'),
}
for name, edit in variants.items():
    mutated = mutate(name + '.xlsx', part, edit)
    assert zipfile.ZipFile(mutated).read(part) != zipfile.ZipFile(fixture).read(part), name
    refusal(name, workbook=mutated)

rels = 'xl/worksheets/_rels/sheet1.xml.rels'
external = mutate('external.xlsx', rels, lambda b: b.replace(b'Target=', b'TargetMode="External" Target='))
refusal('external', workbook=external)
missing = mutate('missing-part.xlsx', rels, lambda b: b.replace(b'table1.xml', b'absent.xml'))
refusal('missing-part', workbook=missing)
worksheet = 'xl/worksheets/sheet1.xml'
formula = mutate('formula.xlsx', worksheet, lambda b: b.replace(b'<c r="D4"><v>12.5</v></c>', b'<c r="D4"><f>1+1</f><v>2</v></c>'))
refusal('formula', workbook=formula)
merged = mutate('merged.xlsx', worksheet, lambda b: b.replace(b'</sheetData>', b'</sheetData><mergeCells count="1"><mergeCell ref="B4:C4"/></mergeCells>'))
refusal('merged', workbook=merged)

# Current definition, not remembered A1 coordinates, controls this import.
shrunk = mutate('smaller-table.xlsx', part, lambda b: b.replace(b'B3:F7', b'B3:F5').replace(b'totalsRowCount="1"', b'totalsRowCount="0"'))
report, data = export('smaller.json', recipe='-', workbook=shrunk)
assert report['worksheet_rows'] == [4, 5]
assert [r['id'] for r in json.loads(data.read_text())] == ['001', '002']
empty = mutate('empty-table.xlsx', part, lambda b: b.replace(b'B3:F7', b'B3:F3').replace(b'totalsRowCount="1"', b'totalsRowCount="0"'))
report, data = export('empty.xlsx', 'xlsx', recipe='-', workbook=empty)
assert report['output_rows'] == 0 and report['worksheet_rows'] == []
assert list(openpyxl.load_workbook(data).active.values) == [('id', 'customer', 'amount', 'active', 'note')]

existing = out / 'orders.json'
before = existing.read_bytes()
call('export', fixture, 'Orders', query, 'json', existing, good=False)
assert existing.read_bytes() == before
call('export', fixture, 'Orders', query, 'xlsx', fixture, good=False)
assert fixture.read_bytes() == original
refusal('directory-input', workbook=out)

# Exact fresh retry after a refusal.
report, data = export('retry.json')
assert json.loads(data.read_text()) == expected
assert fixture.read_bytes() == original
# A non-cell chartsheet must not hide unrelated worksheet Tables.
with_chart = source / 'kit/xlsx-named-table/orders-with-chartsheet.xlsx'
chart_catalog = call('inspect', with_chart)
assert chart_catalog['tables'] == catalog['tables']
chart_report, chart_data = export('with-chartsheet.json', workbook=with_chart)
assert json.loads(chart_data.read_text()) == expected
assert chart_report['worksheet_rows'] == [6, 4]
assert openpyxl.load_workbook(with_chart).chartsheets[0].title == 'Cost chart'

summary = {'cli_calls': len(calls), 'passed': len(calls), 'failed': 0,
           'source_sha256': digest(original), 'producer': 'XlsxWriter 3.2.9',
           'independent_reader': 'openpyxl ' + openpyxl.__version__,
           'checks': ['named definitions across sheets', 'headerless table', 'totals excluded',
                      'hidden rows included', 'unrelated chartsheets skipped', 'query row provenance', 'CSV/JSON/XLSX values and types',
                      'literal formula-looking text', 'empty table schema', 'current ref changes',
                      'refusals create no output', 'fresh destinations', 'unchanged source', 'retry']}
(out / 'summary.json').write_text(json.dumps(summary, indent=2) + '\n')
print(json.dumps(summary, indent=2))
