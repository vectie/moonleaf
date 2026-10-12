#!/usr/bin/env python3
"""Exercise the actual DOCX target CLI; keep all bytes, reports and refusals."""
import copy
import hashlib
import json
import pathlib
import subprocess
import sys
import zipfile
import xml.etree.ElementTree as ET

build, out = map(pathlib.Path, sys.argv[1:])
out.mkdir(parents=True, exist_ok=False)
cli = build / 'js/debug/build/cmd/docx-edit/docx-edit.js'
ns = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
xmlns = 'http://www.w3.org/XML/1998/namespace'
main = '''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
<w:body><!-- keep <w:t>fake</w:t> exactly --><?keep this?>
<w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr><w:r><w:rPr><w:b/></w:rPr><w:t>Acme</w:t></w:r><w:r><w:t>Acme</w:t><w:tab/><w:t> tail</w:t></w:r></w:p>
<w:p><w:r><w:t data-note="a > b">unchanged&#13;line</w:t></w:r><w:hyperlink><w:r><w:t>Acme link</w:t></w:r></w:hyperlink></w:p>
<w:tbl><w:tblPr/><w:tr><w:tc><w:p><w:r><w:t><![CDATA[Acme]]></w:t></w:r></w:p></w:tc><w:tc><w:p><w:r><w:t/></w:r></w:p></w:tc></w:tr></w:tbl>
<w:sdt><w:sdtContent><w:p><w:r><w:t>Acme wrapped</w:t></w:r></w:p></w:sdtContent></w:sdt>
<w:sectPr><w:pgSz w:w="12240" w:h="15840"/></w:sectPr></w:body></w:document>'''
parts = {
    '[Content_Types].xml': ('<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
      '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
      '<Default Extension="xml" ContentType="application/xml"/>'
      '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>').encode(),
    '_rels/.rels': b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>',
    'word/document.xml': main.encode(),
    'word/header1.xml': f'<w:hdr xmlns:w="{ns}"><w:p><w:r><w:t>Acme header</w:t></w:r></w:p></w:hdr>'.encode(),
    'word/styles.xml': f'<w:styles xmlns:w="{ns}"/>'.encode(),
    'custom/opaque.bin': b'\0\xffuninterpreted\r\n',
}

def package(path, replacements=None):
    content = dict(parts)
    content.update(replacements or {})
    with zipfile.ZipFile(path, 'w', zipfile.ZIP_STORED) as z:
        for name, data in content.items():
            z.writestr(name, data)


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()

calls = []
def run(label, *args, ok=True):
    p = subprocess.run(['node', str(cli), *map(str, args)], capture_output=True)
    (out / f'{label}.stdout').write_bytes(p.stdout)
    (out / f'{label}.stderr').write_bytes(p.stderr)
    calls.append({'label':label,'argv':list(map(str,args)),'exit':p.returncode,'expected_success':ok})
    assert (p.returncode == 0) == ok, (label,p.stdout,p.stderr)
    if not ok:
        assert p.stdout == b'', (label,'unexpected success output')
        return None
    return json.loads(p.stdout)

source = out / 'source.docx'
package(source)
original_hash = digest(source)
listing = run('inspect', 'inspect', source)
assert listing['source_sha256'] == original_hash
assert [t['text'] for t in listing['targets']] == ['Acme','Acme',' tail','unchanged\rline','Acme','']
assert len({t['id'] for t in listing['targets']}) == 6
request = {'version':1,'source_sha256':original_hash,'target':listing['targets'][4]['id'],'expected':'Acme','replacement':'  Client B & <雪😀>\r\n\t '}
request_path = out / 'request.json'
request_path.write_text(json.dumps(request, ensure_ascii=False))
output = out / 'client-copy.docx'
report = run('apply', 'apply', source, request_path, output)
assert report['source_sha256'] == original_hash and report['copy_sha256'] == digest(output)
assert report['changed'] and report['reopened'] and report['readback_equal']
assert digest(source) == original_hash
reopened = run('reinspect', 'inspect', output)
assert [t['id'] for t in reopened['targets']] == [t['id'] for t in listing['targets']]
assert reopened['targets'][4]['text'] == request['replacement']
assert reopened['targets'][:4] == listing['targets'][:4]
assert reopened['targets'][5:] == listing['targets'][5:]
with zipfile.ZipFile(output) as z:
    assert z.testzip() is None
    assert set(z.namelist()) == set(parts)
    for name, data in parts.items():
        if name != 'word/document.xml': assert z.read(name) == data, name
    edited_xml = z.read('word/document.xml')
    # Only the selected element is rewritten; comments, PI, CDATA elsewhere,
    # attributes, original spellings and all other main-part bytes are retained.
    before, after = parts['word/document.xml'].split(b'<w:t><![CDATA[Acme]]></w:t>')
    assert edited_xml.startswith(before) and edited_xml.endswith(after)
    edited = ET.fromstring(edited_xml)
    original = ET.fromstring(parts['word/document.xml'])
    target = original.find(f'{{{ns}}}body/{{{ns}}}tbl/{{{ns}}}tr/{{{ns}}}tc/{{{ns}}}p/{{{ns}}}r/{{{ns}}}t')
    target.text = request['replacement']
    target.set(f'{{{xmlns}}}space', 'preserve')
    assert ET.tostring(edited) == ET.tostring(original)

# Same-value edits preserve the main part byte-for-byte; empty slots keep IDs.
for label, index, replacement in [('noop',0,'Acme'),('fill-empty',5,'new'),('erase',1,''),('limit',0,'x'*32768)]:
    req = dict(request,target=listing['targets'][index]['id'],expected=listing['targets'][index]['text'],replacement=replacement)
    path = out / f'{label}.json'; path.write_text(json.dumps(req))
    dest = out / f'{label}.docx'
    result = run(label,'apply',source,path,dest)
    inspected = run(label+'-reopen','inspect',dest)
    assert inspected['targets'][index]['text'] == replacement
    assert [t['id'] for t in inspected['targets']] == [t['id'] for t in listing['targets']]
    if label == 'noop':
        assert not result['changed']
        with zipfile.ZipFile(dest) as z: assert z.read('word/document.xml') == parts['word/document.xml']

# Every refusal is followed by byte checks; no failure may emit a success report.
negative_requests = {
    'stale-hash':dict(request,source_sha256='0'*64),
    'stale-text':dict(request,expected='different'),
    'stale-noop':dict(request,expected='different',replacement='different'),
    'unknown-target':dict(request,target='body/p:999/r:1/t:1'),
    'unsupported-wrapper':dict(request,target='body/sdt:4/sdtContent:1/p:1/r:1/t:1',expected='Acme wrapped'),
    'version':dict(request,version=2),
    'extra-field':dict(request,extra=True),
    'missing-field':{k:v for k,v in request.items() if k!='expected'},
    'wrong-type':dict(request,replacement=1),
    'nul':dict(request,replacement='a\0b'),
    'forbidden-control':dict(request,replacement='a\x01b'),
    'noncharacter':dict(request,replacement='\uffff'),
    'text-over-limit':dict(request,replacement='x'*32769),
}
for label, req in negative_requests.items():
    path = out / f'{label}.json'; path.write_text(json.dumps(req))
    dest = out / f'{label}.docx'
    run(label,'apply',source,path,dest,ok=False)
    assert not dest.exists()
    assert digest(source) == original_hash
existing_hash = digest(output)
run('existing-output','apply',source,request_path,output,ok=False)
assert digest(output) == existing_hash
run('source-as-output','apply',source,request_path,source,ok=False)
assert digest(source) == original_hash
run('changed-source','apply',output,request_path,out/'changed-source.docx',ok=False)
assert not (out/'changed-source.docx').exists()
for label, xml in {
    'fields':main.replace('<w:p><w:pPr>', '<w:p><w:r><w:fldChar w:fldCharType="begin"/></w:r><w:pPr>',1),
    'wrong-namespace':main.replace(ns,'urn:other'),
    'malformed':main.replace('</w:body>',''),
    'duplicate-body':main.replace('</w:document>','<w:body/></w:document>'),
}.items():
    path = out / f'{label}.docx'; package(path,{'word/document.xml':xml.encode()})
    before = digest(path)
    run(label,'inspect',path,ok=False)
    assert digest(path) == before
invalid_request = out/'invalid-utf8.json'; invalid_request.write_bytes(b'\xff')
run('invalid-utf8','apply',source,invalid_request,out/'invalid-utf8.docx',ok=False)
run('usage',ok=False)
run('directory-source','inspect',out,ok=False)
run('valid-retry','apply',source,request_path,out/'retry.docx')
assert digest(source) == original_hash

# Real library-produced DOCX (not an Office visual-fidelity qualification).
producer = {'status':'not_run','reason':'python-docx unavailable'}
try:
    import docx
except ImportError:
    pass
else:
    p = docx.Document()
    heading = p.add_heading('Client proposal',0)
    para = p.add_paragraph(); para.add_run('Acme').bold = True; para.add_run(' / Acme')
    table = p.add_table(rows=1, cols=2); table.cell(0,0).text='Acme'; table.cell(0,1).text='Retain'
    p.sections[0].header.paragraphs[0].text='Header retained'
    path = out/'python-docx-source.docx'; p.save(path)
    before = digest(path)
    found = run('producer-inspect','inspect',path)
    selected = next(t for t in found['targets'] if '/tbl:' in t['id'] and t['text']=='Acme')
    req = dict(request,source_sha256=before,target=selected['id'],replacement='Client B')
    reqpath=out/'producer-request.json'; reqpath.write_text(json.dumps(req))
    dest=out/'python-docx-copy.docx'; run('producer-apply','apply',path,reqpath,dest)
    loaded=docx.Document(dest)
    assert loaded.tables[0].cell(0,0).text=='Client B'
    assert loaded.tables[0].cell(0,1).text=='Retain'
    assert loaded.paragraphs[1].text=='Acme / Acme' and loaded.paragraphs[1].runs[0].bold
    assert loaded.sections[0].header.paragraphs[0].text=='Header retained'
    assert digest(path)==before
    with zipfile.ZipFile(path) as a, zipfile.ZipFile(dest) as b:
        assert set(a.namelist())==set(b.namelist())
        for name in a.namelist():
            if name!='word/document.xml': assert a.read(name)==b.read(name),name
    producer={'status':'passed','producer':'python-docx','version':docx.__version__,'original_sha256':before,'copy_sha256':digest(dest),'original_unchanged':True}
summary={'status':'passed','invocations':len(calls),'refusals':sum(not c['expected_success'] for c in calls),'source_sha256':original_hash,'original_unchanged':digest(source)==original_hash,'copy_sha256':digest(output),'independent_xml_comparison':True,'selected_element_only_splice':True,'non_document_parts_equal':True,'producer':producer,'calls':calls}
(out/'results.json').write_text(json.dumps(summary,indent=2)+'\n')
print(json.dumps({k:v for k,v in summary.items() if k!='calls'},indent=2))
