#!/usr/bin/env python3
"""Focused real-adapter and independent-copy checks. No fixture regeneration."""
import copy, hashlib, json, os, stat, subprocess, sys, zipfile
import xml.etree.ElementTree as ET
from pathlib import Path
import openpyxl

build, out = map(lambda p: Path(p).resolve(), sys.argv[1:])
out.mkdir(parents=True, exist_ok=False)
source = Path(__file__).resolve().parents[1]
kit = source/'kit/xlsx-edit'
original = kit/'contacts.xlsx'
request = json.loads((kit/'corrections.json').read_text())
cli = build/'js/debug/build/cmd/xlsx-edit/xlsx-edit.js'
sha = lambda p: hashlib.sha256(Path(p).read_bytes()).hexdigest()
original_hash = sha(original)
assert request['source_sha256'] == original_hash
calls = []

def invoke(label, args, ok=True, env=None):
    result = subprocess.run(['node', str(cli), *map(str,args)], capture_output=True, text=True, env=env)
    (out/(label+'.stdout')).write_text(result.stdout)
    (out/(label+'.stderr')).write_text(result.stderr)
    calls.append({'label':label,'argv':['node',str(cli),*map(str,args)],'exit_code':result.returncode})
    (out/'calls.json').write_text(json.dumps(calls,indent=2)+'\n')
    assert (result.returncode == 0) == ok, (label,result.returncode,result.stderr)
    if not ok:
        assert not result.stdout.strip(), (label, result.stdout)
        return None
    return json.loads(result.stdout)

def apply(label, req, dest=None, ok=True, src=original, env=None):
    path = out/(label+'.request.json')
    path.write_text(json.dumps(req,ensure_ascii=False,indent=2)+'\n')
    dest = dest or out/(label+'.xlsx')
    result = invoke(label,['apply',src,path,dest],ok,env)
    if not ok and not dest.exists():
        assert not dest.is_symlink()
    return result, dest

inspect = invoke('inspect-shared',['inspect', original,'Customers','B2','B3','E2'])
assert inspect['schema'] == 'moonleaf.xlsx-text-targets.v1'
assert inspect['source_sha256'] == original_hash
assert [t['text'] for t in inspect['targets']] == ['Ada','Bo','Call Monday']
assert all(t['style_id'] is not None for t in inspect['targets'][:2])
hidden = invoke('inspect-hidden',['inspect',original,'Private notes','A1'])
assert hidden['targets'][0]['text'] == 'Private draft'
report, client = apply('apply-batch',request)
assert report['schema'] == 'moonleaf.xlsx-text-copy.v1'
assert report['requested_cells'] == report['changed_cells'] == 4
assert report['source_sha256'] == original_hash and report['copy_sha256'] == sha(client)
assert report['reopened'] and report['readback_equal']
assert stat.S_IMODE(client.stat().st_mode) == 0o600
assert all(s['changes'] == ['cells'] for s in report['comparison']['sheets'])
assert 'stale' in ' '.join(report['warnings'])
old = openpyxl.load_workbook(original, data_only=False)
new = openpyxl.load_workbook(client, data_only=False)
old_cache = openpyxl.load_workbook(original, data_only=True)
new_cache = openpyxl.load_workbook(client, data_only=True)
assert old.sheetnames == new.sheetnames
for e in request['edits']:
    a,b = old[e['sheet']][e['address']], new[e['sheet']][e['address']]
    assert a.value == e['expected'] and b.value == e['replacement']
    assert b.data_type == 's' and a.style_id == b.style_id
    assert a._style == b._style
assert new['Customers']['E2'].value == '=SUM(1,2)' and new['Customers']['E2'].data_type == 's'
assert new['Customers']['E3'].value == 'Call Monday'
assert new['Customers']['A2'].value == '001'
assert new['Customers']['C2'].value == 12.5
for addr in ('D2','D3'):
    assert old['Customers'][addr].value == new['Customers'][addr].value
    assert old_cache['Customers'][addr].value == new_cache['Customers'][addr].value
assert new['Private notes'].sheet_state == 'hidden'
assert str(old['Customers'].merged_cells) == str(new['Customers'].merged_cells)
ns={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
with zipfile.ZipFile(original) as a, zipfile.ZipFile(client) as b:
    assert set(a.namelist()) == set(b.namelist())
    modified={'xl/worksheets/sheet1.xml':{'B2','B3','E2'},'xl/worksheets/sheet2.xml':{'A1'}}
    for part in a.namelist():
        if part not in modified:
            assert a.read(part) == b.read(part), part
        else:
            ar,br=ET.fromstring(a.read(part)),ET.fromstring(b.read(part))
            for tree in (ar,br):
                for row in tree.findall('s:sheetData/s:row',ns):
                    for cell in list(row):
                        if cell.tag == '{'+ns['s']+'}c' and cell.get('r') in modified[part]:
                            row.remove(cell)
            assert ET.tostring(ar)==ET.tostring(br), part
inline=invoke('inspect-inline',['inspect',client,'Customers','B2','E2'])
assert [t['text'] for t in inline['targets']] == ['Ada Lovelace','=SUM(1,2)']
noop=copy.deepcopy(request)
for e in noop['edits']: e['replacement']=e['expected']
r,p=apply('no-op',noop)
assert r['changed_cells']==0 and r['comparison']['sheets']==[]
with zipfile.ZipFile(original) as a, zipfile.ZipFile(p) as b:
    assert all(a.read(name)==b.read(name) for name in a.namelist())
for label, req in [('stale-digest',{**request,'source_sha256':'0'*64}),('empty-batch',{**request,'edits':[]}),('unknown-field',{**request,'extra':1})]:
    _,p=apply(label,req,ok=False); assert not p.exists()
stale=copy.deepcopy(request); stale['edits'][-1]['expected']='stale'
_,p=apply('stale-final-cell',stale,ok=False); assert not p.exists()
dup=copy.deepcopy(request); dup['edits'].append(dup['edits'][0])
_,p=apply('duplicate-target',dup,ok=False); assert not p.exists()
invalid=copy.deepcopy(request); invalid['edits'][0]['replacement']=42
_,p=apply('nonstring-replacement',invalid,ok=False); assert not p.exists()
for label, address, expected in [('rich','B5','Rich name'),('merged-anchor','A7','Merged banner'),('merged-follower','B7',''),('formula','D2','25'),('array-follower','G2','0'),('number','C2','12.5'),('absent','H3','')]:
    req={**request,'edits':[{'sheet':'Customers','address':address,'expected':expected,'replacement':'replacement'}]}
    _,p=apply('refuse-'+label,req,ok=False); assert not p.exists()
    if label in ('rich','formula'):
        invoke('inspect-refuse-'+label,['inspect',original,'Customers',address],ok=False)
apply('same-source',request,dest=original,ok=False)
existing=out/'existing.xlsx'; existing.write_bytes(b'keep me')
apply('existing-output',request,dest=existing,ok=False); assert existing.read_bytes()==b'keep me'
hard=out/'hardlink.xlsx'; os.link(original,hard)
apply('hardlink-output',request,dest=hard,ok=False)
soft=out/'symlink.xlsx'; soft.symlink_to(original)
apply('symlink-output',request,dest=soft,ok=False)
# Inject a disk-readback mismatch after the real exclusive write, preserving its
# actual output for inspection. This is a host-failure test, not a fake dependency.
failed=out/'readback-failed.xlsx'
preloader=out/'readback-failure.cjs'
preloader.write_text("const fs=require('node:fs');const read=fs.readFileSync;fs.readFileSync=function(p,...a){const b=read.call(this,p,...a);return String(p)===process.env.FAIL_COPY?Buffer.concat([b,Buffer.from('fault')]):b;};\n")
env=dict(os.environ,NODE_OPTIONS='--require='+str(preloader),FAIL_COPY=str(failed))
apply('readback-failure',request,dest=failed,ok=False,env=env)
assert failed.exists()
apply('readback-existing-refused',request,dest=failed,ok=False)
retry, retryfile=apply('fresh-retry',request)
assert retry['readback_equal'] and sha(retryfile)==sha(client)
# XML processors normalize literal CR/CRLF. Verify actual copied strings with
# an independent reader and require character references in the emitted XML.
line_edits = [
    {'sheet':'Customers','address':'B2','expected':'Ada','replacement':'  Ada 雪😀 & <L>\r\nnext\tend  '},
    {'sheet':'Customers','address':'E2','expected':'Call Monday','replacement':'before\rafter'},
    {'sheet':'Customers','address':'E3','expected':'Call Monday','replacement':'literal &#13; remains literal'},
]
line_report,line_copy=apply('independent-line-endings',{**request,'edits':line_edits})
line_book=openpyxl.load_workbook(line_copy)
for e in line_edits:
    assert line_book[e['sheet']][e['address']].value == e['replacement']
    assert line_book[e['sheet']][e['address']].data_type == 's'
with zipfile.ZipFile(line_copy) as z:
    payload=z.read('xl/worksheets/sheet1.xml')
    assert b'\r' not in payload
    assert b'&#13;\nnext' in payload and b'before&#13;after' in payload
    assert b'literal &amp;#13; remains literal' in payload
    tree=ET.fromstring(payload)
    cells={c.get('r'):c for c in tree.findall('s:sheetData/s:row/s:c',ns)}
    for e in line_edits:
        assert cells[e['address']].find('s:is/s:t',ns).text == e['replacement']
line_inspect=invoke('inspect-line-endings',['inspect',line_copy,'Customers','B2','E2','E3'])
assert [t['text'] for t in line_inspect['targets']]==[e['replacement'] for e in line_edits]
assert line_report['changed_cells']==3 and line_report['readback_equal']
assert sha(original)==original_hash and sha(hard)==original_hash
result={'status':'passed','invocations':len(calls),'source_sha256':original_hash,'copy_sha256':sha(client),'openpyxl_version':openpyxl.__version__,'changed_cells':4,'source_unchanged':True,'unrelated_part_bytes_unchanged':True,'untargeted_sheet_xml_equal':True,'formula_expressions_and_caches_unchanged':True,'styles_preserved':True,'copy_readback_failure_and_fresh_retry':True,'independent_crlf_cr_and_literal_entity_round_trip':True}
(out/'RESULT.json').write_text(json.dumps(result,indent=2)+'\n')
print(json.dumps(result,indent=2))
