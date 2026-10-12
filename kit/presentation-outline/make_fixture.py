"""Regenerate the actual python-pptx producer fixture; requires python-pptx 1.0.2."""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches
import pptx
assert pptx.__version__ == '1.0.2', pptx.__version__
out = Path(__file__).parent / 'quarterly-review.pptx'
prs = Presentation()
first = prs.slides.add_slide(prs.slide_layouts[1])
first.shapes.title.text = 'Operations | October'
first.placeholders[1].text = 'Ship the pilot\nReview supplier costs\n雪😀 & partners'
first.notes_slide.notes_text_frame.text = 'Ask for a decision.\nOwner: operations\nDo not distribute the draft.'
second = prs.slides.add_slide(prs.slide_layouts[5])
second.shapes.title.text = 'Budget <draft>'
second._element.set('show', '0')
table = second.shapes.add_table(3, 2, Inches(1), Inches(2), Inches(6), Inches(2)).table
for row, values in enumerate([['Item','USD'],['Hosting','1200'],['Research','Line 1\nLine 2']]):
    for col, value in enumerate(values): table.cell(row, col).text = value
second.notes_slide.notes_text_frame.text = 'Check the research estimate.\nBudget is not approved.'
# Actual presentation order intentionally differs from slide part numbering.
ids = prs.slides._sldIdLst
last = ids[-1]
ids.remove(last)
ids.insert(0, last)
prs.save(out)
print(out)
