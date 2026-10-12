#!/usr/bin/env python3
"""Generate the retained business-table fixture using python-docx 1.2.0.
Not required for tests, which consume the retained DOCX bytes.
"""
from pathlib import Path
from docx import Document
root = Path(__file__).resolve().parent
doc = Document()
doc.add_heading('Supplier comparison', 0)
doc.add_paragraph('Review the second table; the cover table is merged.')
cover = doc.add_table(rows=1, cols=2)
cover.cell(0,0).merge(cover.cell(0,1)).text = 'Draft supplier comparison'
doc.add_paragraph('Prices are source text, not approved commitments.')
table = doc.add_table(rows=1, cols=4)
table.style = 'Table Grid'
for cell, text in zip(table.rows[0].cells, ['id','supplier','amount','notes']):
    cell.text = text
for values in [('001','雪😀 & Co','1200.50','Line 1\nLine 2'), ('002','Beta, "Labs"','900',''), ('003','Gamma','1500','=1+1')]:
    row = table.add_row()
    for cell, text in zip(row.cells, values): cell.text = text
# Split rich text into runs while keeping the exact visible value.
p = table.cell(1,1).paragraphs[0]; p.clear(); p.add_run('雪😀').bold = True; p.add_run(' & Co')
# Multiple paragraphs, including an empty middle paragraph.
cell = table.cell(1,3); cell.text = 'Line 1'; cell.add_paragraph(''); cell.add_paragraph('Line 2\tend')
doc.sections[0].header.paragraphs[0].text = 'Header is not table data'
doc.save(root/'suppliers.docx')
