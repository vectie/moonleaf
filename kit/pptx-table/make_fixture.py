"""Optional fixture regeneration with the actual python-pptx 1.0.2 producer."""
import hashlib,json
from pathlib import Path
import pptx
from pptx import Presentation
from pptx.util import Inches
assert pptx.__version__ == '1.0.2', pptx.__version__
kit=Path(__file__).parent
prs=Presentation()
cover=prs.slides.add_slide(prs.slide_layouts[5]); cover.shapes.title.text='Supplier review'
summary=cover.shapes.add_table(1,2,Inches(1),Inches(2),Inches(7),Inches(1)).table
summary.cell(0,0).text='Code';summary.cell(0,1).text='Description'
slide=prs.slides.add_slide(prs.slide_layouts[5]);slide.shapes.title.text='Source costs';slide._element.set('show','0')
merged=slide.shapes.add_table(1,2,Inches(1),Inches(1),Inches(7),Inches(.5)).table
merged.cell(0,0).merge(merged.cell(0,1));merged.cell(0,0).text='Draft estimates'
shape=slide.shapes.add_table(4,4,Inches(1),Inches(2),Inches(7),Inches(3))
group=slide.shapes.add_group_shape([shape])
table=shape.table
rows=[['id','supplier','amount','notes'],['001','雪😀 & Co','1200.50','Line 1\n\nLine 2\vend'],['002','Beta, "Labs"','900',''],['003','Gamma','1500','=1+1']]
for r,values in enumerate(rows):
    for c,value in enumerate(values):table.cell(r,c).text=value
slide.notes_slide.notes_text_frame.text='Private notes are not table data.'
ids=prs.slides._sldIdLst;last=ids[-1];ids.remove(last);ids.insert(0,last)
output=kit/'suppliers.pptx';prs.save(output)
selection={'version':1,'source_sha256':hashlib.sha256(output.read_bytes()).hexdigest(),'table':'slide:1/table:2','header':'first_row'}
(kit/'selection.json').write_text(json.dumps(selection,indent=2)+'\n')
print(output)
