"""Regenerate with python-pptx 1.0.2; ZIP timestamps may change the digest."""
from pathlib import Path
from io import BytesIO
import base64
import pptx
from pptx.util import Inches
assert pptx.__version__ == '1.0.2'
p = pptx.Presentation()
for number in (1, 2):
    s = p.slides.add_slide(p.slide_layouts[6])
    title = s.shapes.add_textbox(Inches(1), Inches(.5), Inches(5), Inches(.5))
    title.text = f'Client review {number}'
    shape = s.shapes.add_textbox(Inches(1), Inches(1.5), Inches(5), Inches(.5))
    para = shape.text_frame.paragraphs[0]
    r = para.add_run(); r.text = 'Acme'; r.font.bold = True
    r = para.add_run(); r.text = 'Acme'
    group = s.shapes.add_group_shape()
    grouped = group.shapes.add_textbox(Inches(1), Inches(2.5), Inches(5), Inches(.5))
    grouped.text = 'Acme'
    table = s.shapes.add_table(1, 2, Inches(1), Inches(3.5), Inches(5), Inches(.6)).table
    table.cell(0, 0).text = 'Acme table'; table.cell(0, 1).text = 'Retain'
    s.notes_slide.notes_text_frame.text = 'Acme speaker note'
    png = base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+jRZkAAAAASUVORK5CYII=')
    s.shapes.add_picture(BytesIO(png), Inches(6), Inches(1), width=Inches(.3))
    if number == 2: s._element.set('show', '0')
# Human slide order differs from filenames.
ids = p.slides._sldIdLst
ids.insert(0, ids[-1])
p.save(Path(__file__).with_name('source.pptx'))
